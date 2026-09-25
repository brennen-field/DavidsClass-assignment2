from pathlib import Path

import fitz
import pytest
from docx import Document as DocxDocument
from pptx import Presentation
from pptx.util import Inches

from course_assistant.documents import parsing
from course_assistant.documents.convert import find_soffice
from course_assistant.documents.store import DocumentStore

needs_libreoffice = pytest.mark.skipif(find_soffice() is None, reason="LibreOffice not installed")


def make_pdf(path: Path, page_texts: list[str]) -> Path:
    pdf = fitz.open()
    for text in page_texts:
        pdf.new_page().insert_text((72, 72), text)
    pdf.save(path)
    return path


def make_pptx(path: Path, slide_texts: list[str], hidden: set[int] = frozenset()) -> Path:
    deck = Presentation()
    for i, text in enumerate(slide_texts, start=1):
        slide = deck.slides.add_slide(deck.slide_layouts[5])
        slide.shapes.title.text = text
        if i in hidden:
            slide._element.set("show", "0")
    deck.save(path)
    return path


@pytest.fixture
def store(tmp_path):
    return DocumentStore(tmp_path / "data")


@pytest.fixture
def syllabus(tmp_path):
    return make_pdf(tmp_path / "Syllabus.pdf", ["Grading policy", "Late work", "Office hours"])


# ---- adding ----


def test_add_pdf_saves_page_images_and_text(store, syllabus):
    result = store.add_file(syllabus)

    assert result.status == "added"
    pages = store.get_pages(result.document.doc_id)
    assert [p.page_number for p in pages] == [1, 2, 3]
    assert pages[0].doc_name == "Syllabus.pdf"
    assert "Grading policy" in pages[0].text
    assert all(Path(p.image_path).is_file() for p in pages)


def test_same_file_twice_is_not_duplicated(store, syllabus):
    first = store.add_file(syllabus)
    second = store.add_file(syllabus)

    assert second.status == "duplicate"
    assert second.document.doc_id == first.document.doc_id
    assert len(store.list_documents()) == 1


def test_same_content_under_new_name_is_duplicate(store, syllabus, tmp_path):
    store.add_file(syllabus)
    copy = tmp_path / "Syllabus (1).pdf"
    copy.write_bytes(syllabus.read_bytes())

    assert store.add_file(copy).status == "duplicate"
    assert len(store.list_documents()) == 1


def test_different_file_with_same_name_is_rejected(store, syllabus, tmp_path):
    store.add_file(syllabus)
    other = make_pdf(tmp_path / "other.pdf", ["Updated syllabus"])

    result = store.add_file(other, doc_name="Syllabus.pdf")

    assert result.status == "error"
    assert len(store.list_documents()) == 1


def test_unsupported_file_type_is_rejected(store, tmp_path):
    notes = tmp_path / "notes.xlsx"
    notes.write_bytes(b"not really a spreadsheet")

    result = store.add_file(notes)

    assert result.status == "error"
    assert "unsupported" in result.message
    assert store.list_documents() == []


def test_empty_file_is_rejected(store, tmp_path):
    empty = tmp_path / "empty.pdf"
    empty.write_bytes(b"")

    assert store.add_file(empty).status == "error"


def test_corrupt_pdf_leaves_nothing_behind(store, tmp_path):
    broken = tmp_path / "broken.pdf"
    broken.write_bytes(b"%PDF-1.4 this is not a real pdf")

    result = store.add_file(broken)

    assert result.status == "error"
    assert store.list_documents() == []
    assert list(store.docs_dir.iterdir()) == []


def test_documents_survive_restart(store, syllabus):
    doc_id = store.add_file(syllabus).document.doc_id

    reopened = DocumentStore(store.data_dir)

    assert [d.doc_id for d in reopened.list_documents()] == [doc_id]
    assert len(reopened.get_pages(doc_id)) == 3


# ---- parsing service ----


def test_parsing_service_down_falls_back_to_local(store, syllabus, monkeypatch):
    def service_down(pdf_path):
        raise parsing.ServiceUnavailable("connection refused")

    monkeypatch.setattr(parsing, "parse_with_class_service", service_down)

    result = store.add_file(syllabus)

    assert result.status == "added"
    assert result.document.text_source == "local"
    assert "Grading policy" in store.get_pages(result.document.doc_id)[0].text


def test_parsing_service_text_is_used_when_available(store, syllabus, monkeypatch):
    monkeypatch.setattr(parsing, "parse_with_class_service", lambda p: ["a", "b", "c"])

    result = store.add_file(syllabus)

    assert result.document.text_source == "class-service"
    assert [p.text for p in store.get_pages(result.document.doc_id)] == ["a", "b", "c"]


def test_parsing_service_wrong_page_count_falls_back(store, syllabus, monkeypatch):
    monkeypatch.setattr(parsing, "parse_with_class_service", lambda p: ["only one"])

    result = store.add_file(syllabus)

    assert result.document.text_source == "local"


# ---- hooks for search indexes ----


def test_listeners_receive_pages_on_add(store, syllabus):
    received = []
    store.on_document_added(received.append)

    store.add_file(syllabus)

    assert len(received) == 1
    assert [p.page_number for p in received[0]] == [1, 2, 3]


def test_failing_listener_keeps_document_and_warns(store, syllabus):
    def broken_indexer(pages):
        raise RuntimeError("embedding service down")

    store.on_document_added(broken_indexer)

    result = store.add_file(syllabus)

    assert result.status == "added"
    assert result.warnings
    assert len(store.list_documents()) == 1


# ---- removing ----


def test_remove_deletes_files_catalog_and_index_entries(store, syllabus):
    deleted_from_text, deleted_from_visual = [], []
    store.register_index_deleter(deleted_from_text.append)
    store.register_index_deleter(deleted_from_visual.append)
    doc_id = store.add_file(syllabus).document.doc_id

    result = store.remove(doc_id)

    assert result.ok
    assert deleted_from_text == [doc_id]
    assert deleted_from_visual == [doc_id]
    assert store.list_documents() == []
    assert store.get_pages(doc_id) == []
    assert not (store.docs_dir / doc_id).exists()


def test_failed_index_delete_keeps_document_for_retry(store, syllabus):
    def broken_deleter(doc_id):
        raise RuntimeError("index locked")

    store.register_index_deleter(broken_deleter)
    doc_id = store.add_file(syllabus).document.doc_id

    result = store.remove(doc_id)

    assert not result.ok
    assert store.get_document(doc_id) is not None
    assert len(store.get_pages(doc_id)) == 3


def test_removed_file_can_be_added_again(store, syllabus):
    doc_id = store.add_file(syllabus).document.doc_id
    store.remove(doc_id)

    assert store.add_file(syllabus).status == "added"


def test_remove_unknown_document(store):
    assert not store.remove("does-not-exist").ok


# ---- PowerPoint / Word conversion ----


@needs_libreoffice
def test_pptx_upload_has_one_page_per_slide(store, tmp_path):
    deck = make_pptx(tmp_path / "Week 2.pptx", ["Intro", "Vibe coding in Prod", "Wrap-up"])

    result = store.add_file(deck)

    assert result.status == "added", result.message
    pages = store.get_pages(result.document.doc_id)
    assert len(pages) == 3
    assert "Vibe coding in Prod" in pages[1].text
    assert pages[1].source_format == "pptx"


@needs_libreoffice
def test_hidden_slides_keep_slide_numbers_aligned(store, tmp_path):
    deck = make_pptx(tmp_path / "Deck.pptx", ["One", "Two (hidden)", "Three"], hidden={2})

    pages = store.get_pages(store.add_file(deck).document.doc_id)

    assert len(pages) == 3
    assert "Three" in pages[2].text


@needs_libreoffice
def test_docx_upload(store, tmp_path):
    doc = DocxDocument()
    doc.add_paragraph("Course policies")
    doc.save(tmp_path / "Policies.docx")

    result = store.add_file(tmp_path / "Policies.docx")

    assert result.status == "added", result.message
    assert "Course policies" in store.get_pages(result.document.doc_id)[0].text
